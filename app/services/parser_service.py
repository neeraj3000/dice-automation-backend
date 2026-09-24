import re
from pathlib import Path
from typing import Dict, List, Any

# Common role patterns for heuristic extraction
COMMON_ROLES = [
    "Databricks Consultant", "Databricks Engineer", "Data Engineer", "Senior Data Engineer",
    "Lead Data Engineer", "AI Engineer", "GenAI Engineer", "LLM Engineer", "Machine Learning Engineer",
    "MLOps Engineer", "Data Scientist", "Full Stack Developer", "Backend Engineer", "Frontend Engineer",
    "Software Engineer", "Cloud Architect", "Azure Solutions Architect", "AWS Solutions Architect",
    "DevOps Engineer", "Site Reliability Engineer", "Database Administrator", "Business Intelligence Engineer",
    "Data Analyst", "Product Manager", "Scrum Master"
]

# Common technical skills to scan for
SKILL_KEYWORDS = [
    "Python", "SQL", "PySpark", "Databricks", "Delta Lake", "Azure", "ADF", "Azure Data Factory",
    "Synapse", "AWS", "GCP", "Google Cloud", "Snowflake", "Kafka", "Docker", "Kubernetes",
    "Terraform", "CI/CD", "Git", "FastAPI", "Flask", "Django", "React", "TypeScript",
    "JavaScript", "Node.js", "Next.js", "MongoDB", "PostgreSQL", "MySQL", "Redis",
    "LangChain", "LangGraph", "LlamaIndex", "OpenAI", "RAG", "Vector DB", "Pinecone",
    "Machine Learning", "Deep Learning", "NLP", "Computer Vision", "PyTorch", "TensorFlow",
    "Scikit-Learn", "Pandas", "NumPy", "Hadoop", "Spark", "Airflow", "dbt", "Power BI",
    "Tableau", "Linux", "REST API", "GraphQL", "Microservices", "Java", "Spring Boot",
    "C++", "C#", ".NET", "Go", "Golang", "Rust"
]

def extract_text_from_pdf(file_path: Path) -> str:
    import pymupdf
    text = []
    with pymupdf.open(file_path) as doc:
        for page in doc:
            page_text = page.get_text()
            if page_text:
                text.append(page_text)
    return "\n".join(text).strip()

def extract_text_from_docx(file_path: Path) -> str:
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paragraphs.append(row_text)
    return "\n".join(paragraphs).strip()

def extract_text(file_path: Path, file_type: str) -> str:
    ext = file_type.lower().replace(".", "")
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    elif ext in ["docx", "doc"]:
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {file_type}")

def extract_heuristic_metadata(text: str, file_name: str) -> Dict[str, Any]:
    combined_corpus = f"{file_name}\n{text}"
    
    # 1. Detect target role
    detected_role = "Software Professional"
    for role in COMMON_ROLES:
        if re.search(r"\b" + re.escape(role) + r"\b", combined_corpus, re.IGNORECASE):
            detected_role = role
            break
            
    # 2. Detect skills
    detected_skills: List[str] = []
    for skill in SKILL_KEYWORDS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text, re.IGNORECASE):
            detected_skills.append(skill)
            
    # 3. Detect experience
    detected_exp = "3+ years"
    exp_matches = re.findall(r"(\d{1,2}\+?)\s*(?:years|yrs)\s*(?:of\s*)?(?:experience)?", text, re.IGNORECASE)
    if exp_matches:
        detected_exp = f"{exp_matches[0]} years"
        
    # 4. Create concise summary
    lines = [line.strip() for line in text.split("\n") if len(line.strip()) > 30]
    summary = ""
    for line in lines[:5]:
        if not re.search(r"@|linkedin|github|phone|email|\d{3}-\d{3}", line, re.IGNORECASE):
            summary = line
            break
    if not summary:
        summary = f"{detected_role} with {detected_exp} experience in {', '.join(detected_skills[:5])}."
    elif len(summary) > 250:
        summary = summary[:250] + "..."

    return {
        "target_role": detected_role,
        "skills": detected_skills,
        "experience_years": detected_exp,
        "summary": summary
    }
