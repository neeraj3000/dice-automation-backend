import sys
import asyncio
import io
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pymupdf
import docx
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.database import connect_db, close_db

def create_sample_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        "John Doe\n"
        "Databricks Consultant Resume\n"
        "10+ years of experience in Databricks, PySpark, Azure, ADF, Spark, SQL, Delta Lake.\n"
        "Summary: Senior Databricks Consultant with over 10 years experience architecting lakehouses.\n"
    )
    page.insert_text((50, 50), text)
    return doc.tobytes()

def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Jane Smith - AI Engineer", level=1)
    doc.add_paragraph("AI Engineer with 5 years of experience in Python, FastAPI, LangChain, RAG, Docker, Kubernetes.")
    doc.add_paragraph("Summary: Experienced AI Engineer specializing in LLM systems and production RAG.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

async def run_all_tests():
    print("Connecting to MongoDB...")
    await connect_db()
    
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Test Health
        health = await client.get("/api/health")
        print(f"Health check: {health.status_code} -> {health.json()}")
        assert health.status_code == 200

        # 1. Upload PDF
        print("Testing PDF upload...")
        pdf_bytes = create_sample_pdf_bytes()
        files = [
            ("files", ("Databricks_Consultant.pdf", pdf_bytes, "application/pdf"))
        ]
        res = await client.post("/api/resumes/upload", files=files)
        assert res.status_code == 200, res.text
        pdf_resume = res.json()[0]
        print(f"Uploaded PDF parsed: role='{pdf_resume['target_role']}', skills={pdf_resume['skills']}")
        assert "Databricks" in pdf_resume["target_role"] or "Databricks" in pdf_resume["display_name"]
        assert "Databricks" in pdf_resume["skills"]
        pdf_id = pdf_resume["id"]

        # 2. Upload DOCX
        print("Testing DOCX upload...")
        docx_bytes = create_sample_docx_bytes()
        files = [
            ("files", ("AI_Engineer_v1.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
        res2 = await client.post("/api/resumes/upload", files=files)
        assert res2.status_code == 200, res2.text
        docx_resume = res2.json()[0]
        print(f"Uploaded DOCX parsed: role='{docx_resume['target_role']}', skills={docx_resume['skills']}")
        assert "AI Engineer" in docx_resume["target_role"] or "AI Engineer" in docx_resume["display_name"]
        assert "Python" in docx_resume["skills"]
        docx_id = docx_resume["id"]

        # 3. List Resumes
        print("Testing List Resumes...")
        res_list = await client.get("/api/resumes")
        assert res_list.status_code == 200
        all_resumes = res_list.json()
        ids = [r["id"] for r in all_resumes]
        assert pdf_id in ids
        assert docx_id in ids
        print(f"Total resumes retrieved: {len(all_resumes)}")

        # 4. Search
        print("Testing Search...")
        search_res = await client.get("/api/resumes?search=Databricks")
        assert search_res.status_code == 200
        assert any(r["id"] == pdf_id for r in search_res.json())

        # 5. Detail
        print("Testing Get Detail...")
        detail_res = await client.get(f"/api/resumes/{pdf_id}")
        assert detail_res.status_code == 200
        assert "Databricks" in detail_res.json()["raw_text"]

        # 6. Update with Custom Field (MongoDB flex schema verification)
        print("Testing Update with dynamic custom fields...")
        update_data = {
            "display_name": "Senior Databricks Principal",
            "custom_fields": {"hourly_rate": 120, "security_clearance": "Public Trust"}
        }
        put_res = await client.put(f"/api/resumes/{pdf_id}", json=update_data)
        assert put_res.status_code == 200
        updated = put_res.json()
        assert updated["display_name"] == "Senior Databricks Principal"
        assert updated["custom_fields"]["hourly_rate"] == 120
        print("Dynamic custom fields successfully persisted in MongoDB!")

        # 7. Delete
        print("Testing Deletion...")
        del1 = await client.delete(f"/api/resumes/{pdf_id}")
        assert del1.status_code == 200
        del2 = await client.delete(f"/api/resumes/{docx_id}")
        assert del2.status_code == 200

        get_del = await client.get(f"/api/resumes/{pdf_id}")
        assert get_del.status_code == 404
        print("Resumes deleted cleanly.")

    await close_db()
    print("ALL TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    asyncio.run(run_all_tests())
