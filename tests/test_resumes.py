import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pytest
import io
import pymupdf
import docx
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.database import connect_db, close_db, get_database

@pytest.fixture(autouse=True)
async def setup_db():
    await connect_db()
    yield
    await close_db()

def create_sample_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    text = (
        "John Doe\n"
        "Databricks Consultant Resume\n"
        "10+ years of experience in Databricks, PySpark, Azure, ADF, Spark, SQL, Delta Lake.\n"
        "Summary: Senior Databricks Consultant with over 10 years experience architecting lakehouses.\n"
    )
    page.insert_text((50, 50), text)
    return doc.tobytes()

def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Jane Smith - AI Engineer", level=1)
    doc.add_paragraph("AI Engineer with 5 years of experience in Python, FastAPI, LangChain, RAG, Docker, Kubernetes.")
    doc.add_paragraph("Summary: Experienced AI Engineer specializing in LLM systems and production RAG.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.mark.anyio
async def test_resume_lifecycle():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Upload PDF
        pdf_bytes = create_sample_pdf_bytes()
        files = [
            ("files", ("Databricks_Consultant.pdf", pdf_bytes, "application/pdf"))
        ]
        res = await client.post("/api/resumes/upload", files=files)
        assert res.status_code == 200, res.text
        uploaded = res.json()
        assert len(uploaded) == 1
        pdf_resume = uploaded[0]
        assert "Databricks" in pdf_resume["target_role"] or "Databricks" in pdf_resume["display_name"]
        assert "Databricks" in pdf_resume["skills"]
        assert "PySpark" in pdf_resume["skills"]
        pdf_id = pdf_resume["id"]

        # 2. Upload DOCX
        docx_bytes = create_sample_docx_bytes()
        files = [
            ("files", ("AI_Engineer_v1.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
        ]
        res2 = await client.post("/api/resumes/upload", files=files)
        assert res2.status_code == 200, res2.text
        docx_resume = res2.json()[0]
        assert "AI Engineer" in docx_resume["target_role"] or "AI Engineer" in docx_resume["display_name"]
        assert "Python" in docx_resume["skills"]
        docx_id = docx_resume["id"]

        # 3. List Resumes
        res_list = await client.get("/api/resumes")
        assert res_list.status_code == 200
        all_resumes = res_list.json()
        ids = [r["id"] for r in all_resumes]
        assert pdf_id in ids
        assert docx_id in ids

        # 4. Search Resumes
        search_res = await client.get("/api/resumes?search=Databricks")
        assert search_res.status_code == 200
        matched = search_res.json()
        assert any(r["id"] == pdf_id for r in matched)

        # 5. Get Detail
        detail_res = await client.get(f"/api/resumes/{pdf_id}")
        assert detail_res.status_code == 200
        detail = detail_res.json()
        assert "Databricks" in detail["raw_text"]

        # 6. Update Resume & Custom Field
        update_data = {
            "display_name": "Senior Databricks Principal",
            "custom_fields": {"hourly_rate": 120, "security_clearance": "Public Trust"}
        }
        put_res = await client.put(f"/api/resumes/{pdf_id}", json=update_data)
        assert put_res.status_code == 200
        updated = put_res.json()
        assert updated["display_name"] == "Senior Databricks Principal"
        assert updated["custom_fields"]["hourly_rate"] == 120

        # 6b. Replace Resume File
        replacement_pdf = create_sample_pdf_bytes()
        replace_files = [
            ("file", ("Databricks_Updated_v2.pdf", replacement_pdf, "application/pdf"))
        ]
        rep_res = await client.post(f"/api/resumes/{pdf_id}/replace", files=replace_files)
        assert rep_res.status_code == 200, rep_res.text
        replaced = rep_res.json()
        assert replaced["id"] == pdf_id
        assert "Databricks_Updated_v2" in replaced["file_name"]

        # 7. Delete Resumes
        del1 = await client.delete(f"/api/resumes/{pdf_id}")
        assert del1.status_code == 200
        assert del1.json()["success"] is True

        del2 = await client.delete(f"/api/resumes/{docx_id}")
        assert del2.status_code == 200
        assert del2.json()["success"] is True

        # Verify deletion
        get_deleted = await client.get(f"/api/resumes/{pdf_id}")
        assert get_deleted.status_code == 404
